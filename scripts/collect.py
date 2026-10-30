"""Сборщик реальных данных Росстата: цены производителей и еженедельная цена дизтоплива по РФ.

Запуск (из корня проекта):
  .venv/bin/python scripts/collect.py            # скачать свежие файлы с rosstat.gov.ru и пересобрать ряды
  .venv/bin/python scripts/collect.py --offline  # без сети: разобрать файлы, уже лежащие в raw/

Что делает
 1. Скачивает страницу https://rosstat.gov.ru/statistics/price (проверка TLS включена; сертификаты Минцифры —
    scripts/certs/russian_ca_bundle.pem; verify=False не используется).
 2. Находит ссылки cena_sx_MM-YYYY.xlsx (имя меняется ежемесячно) и nedel_sred_cen.xlsx, качает в raw/.
    Старые файлы не удаляются; прежний nedel_sred_cen.xlsx при изменении содержимого сохраняется под датой.
 3. Разбирает файлы и пишет ряды в public/data (demo: false):
      producer_price_{регион}_{культура}_month   — субъекты из src/config/regions.ts, только реальные месяцы (с 2025);
      producer_price_russia_{культура}_month     — РФ помесячно с 2021;
      diesel_price_russia_week                   — дизтопливо по РФ, понедельно с 2022.
 4. Пересобирает index.json и запускает validate_data.py.

Загрузка идёт через urllib (ssl): requests/urllib3 на системном LibreSSL 2.8 (macOS) падает на getpeercert()
для цепочки Минцифры, а stdlib проверяет цепочку и имя хоста полностью.
"""
import argparse
import json
import re
import shutil
import ssl
import subprocess
import sys
import urllib.error
import urllib.parse
import urllib.request
from datetime import date, datetime, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import DATA_DIR, RAW_DIR, ROOT, build_manifest, make_series, validate_series, write_series  # noqa: E402

PAGE_URL = "https://rosstat.gov.ru/statistics/price"
CA_BUNDLE = ROOT / "scripts" / "certs" / "russian_ca_bundle.pem"
META = RAW_DIR / "collect_meta.json"

CROP_CODES = {"wheat": "01.11.1", "barley": "01.11.31", "corn": "01.11.20", "sunflower": "01.11.95"}  # ОКПД2
MONTHS_NOM = ["январь", "февраль", "март", "апрель", "май", "июнь", "июль", "август", "сентябрь", "октябрь", "ноябрь", "декабрь"]
MONTHS_GEN = ["января", "февраля", "марта", "апреля", "мая", "июня", "июля", "августа", "сентября", "октября", "ноября", "декабря"]
RF = "Российская Федерация"

PRICE_SOURCE = ("Росстат: Средние цены производителей сельскохозяйственной продукции, реализованной "
                "сельскохозяйственными организациями (на конец периода)")
DIESEL_SOURCE = "Росстат: Еженедельные средние потребительские цены на отдельные товары и услуги по Российской Федерации"


class CollectError(Exception):
    pass


def norm(s):
    return " ".join(str(s).split()) if s is not None else ""


# ── скачивание ────────────────────────────────────────────────────────────────
def http_get(url, timeout=120):
    if not CA_BUNDLE.exists():
        raise CollectError("Нет файла сертификатов {}".format(CA_BUNDLE))
    ctx = ssl.create_default_context(cafile=str(CA_BUNDLE))  # проверка цепочки и имени хоста включена
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 (agromonitoring collector)"})
    try:
        with urllib.request.urlopen(req, context=ctx, timeout=timeout) as r:
            return r.read()
    except (urllib.error.URLError, OSError) as e:
        raise CollectError("Не удалось скачать {}: {}. Проверьте сеть или используйте --offline.".format(url, e))


def find_links(html):
    cena, nedel = [], None
    for m in re.finditer(r'href="([^"]*?(cena_sx_(\d{2})-(\d{4})\.xlsx|nedel_sred_cen\.xlsx))"', html):
        url = urllib.parse.urljoin(PAGE_URL, m.group(1))
        if m.group(2).startswith("cena_sx"):
            cena.append(((int(m.group(4)), int(m.group(3))), url, m.group(2)))
        else:
            nedel = url
    if not cena or not nedel:
        raise CollectError("На странице {} не найдены ссылки на cena_sx_MM-YYYY.xlsx и/или nedel_sred_cen.xlsx "
                           "(вёрстка сайта изменилась?).".format(PAGE_URL))
    cena.sort()
    return cena[-1][1], cena[-1][2], nedel


def download():
    html = http_get(PAGE_URL).decode("utf-8", "replace")
    cena_url, cena_name, nedel_url = find_links(html)
    RAW_DIR.mkdir(exist_ok=True)
    meta = json.loads(META.read_text(encoding="utf-8")) if META.exists() else {}
    cena_path = RAW_DIR / cena_name
    cena_path.write_bytes(http_get(cena_url))
    nedel_path = RAW_DIR / "nedel_sred_cen.xlsx"
    new = http_get(nedel_url)
    if nedel_path.exists() and nedel_path.read_bytes() != new:  # старую версию не теряем
        stamp = datetime.fromtimestamp(nedel_path.stat().st_mtime).strftime("%Y-%m-%d")
        shutil.copy2(nedel_path, RAW_DIR / "nedel_sred_cen_{}.xlsx".format(stamp))
    nedel_path.write_bytes(new)
    meta[cena_name] = cena_url
    meta["nedel_sred_cen.xlsx"] = nedel_url
    META.write_text(json.dumps(meta, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print("скачано:", cena_name, "и nedel_sred_cen.xlsx")
    return cena_path, nedel_path


def latest_local():
    files = []
    for p in RAW_DIR.glob("cena_sx_*.xlsx"):
        m = re.fullmatch(r"cena_sx_(\d{2})-(\d{4})\.xlsx", p.name)
        if m:
            files.append(((int(m.group(2)), int(m.group(1))), p))
    nedel = RAW_DIR / "nedel_sred_cen.xlsx"
    if not files or not nedel.exists():
        raise CollectError("В raw/ нет cena_sx_MM-YYYY.xlsx и/или nedel_sred_cen.xlsx. Запустите без --offline.")
    return sorted(files)[-1][1], nedel


def url_for(name, default):
    if META.exists():
        return json.loads(META.read_text(encoding="utf-8")).get(name, default)
    return default


# ── разбор ────────────────────────────────────────────────────────────────────
def cell_value(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = norm(v).replace(",", ".").replace(" ", "")
    if s in ("", "-", "—", "...", "…", "x"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def region_names():
    """{название субъекта: id} из src/config/regions.ts — новые регионы подхватываются автоматически."""
    txt = (ROOT / "src" / "config" / "regions.ts").read_text(encoding="utf-8")
    res = dict(re.findall(r"id: '(\w+)',\s*name: '([^']+)'", txt))
    return {name: rid for rid, name in res.items()}


def parse_prices(path):
    """-> {(scope, crop): {'YYYY-MM-01': value|None}}, updated_at 'YYYY-MM-DD'."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    names = region_names()
    wanted = dict(names, **{RF: "russia"})
    out = {}
    updated = None
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if ws.title.lower().startswith("содерж"):
            for r in rows:
                m = re.search(r"Обновлено:\s*(\d{2})\.(\d{2})\.(\d{4})", " ".join(norm(c) for c in r if c))
                if m:
                    updated = "{}-{}-{}".format(m.group(3), m.group(2), m.group(1))
            continue
        head = " ".join(norm(c) for r in rows[:4] for c in r if c)
        ym = re.search(r"\bв (\d{4}) г", head)
        if not ym or "на конец периода" not in head:
            continue  # 1.1/1.2, «в среднем за год» (6.2) и т. п.
        hdr = next((r for r in rows[:8] if any(norm(c).lower() == "январь" for c in r)), None)
        if hdr is None:
            continue
        year = int(ym.group(1))
        cols = {MONTHS_NOM.index(norm(c).lower()) + 1: i for i, c in enumerate(hdr) if norm(c).lower() in MONTHS_NOM}
        regional = "субъектам" in head
        code_to_crop = {v: k for k, v in CROP_CODES.items()}
        cur = None
        for r in rows:
            code = norm(r[1]) if len(r) > 1 else ""
            name = norm(r[0])
            if not regional:
                if code in code_to_crop:
                    key = ("russia", code_to_crop[code])
                    out.setdefault(key, {}).update(
                        {"{}-{:02d}-01".format(year, m): cell_value(r[i]) for m, i in cols.items()})
                continue
            if code:  # строка-заголовок культуры
                cur = code_to_crop.get(code) if code in code_to_crop else None
                continue
            if cur and name in wanted:
                key = (wanted[name], cur)
                pts = {"{}-{:02d}-01".format(year, m): cell_value(r[i]) for m, i in cols.items()}
                if not any(k.startswith(str(year)) for k in out.get(key, {})):  # первое вхождение в листе
                    out.setdefault(key, {}).update(pts)
    if not out:
        raise CollectError("В файле {} не найдено ни одной таблицы цен производителей.".format(path.name))
    return out, updated


def parse_diesel(path):
    """-> ({'YYYY-MM-DD' понедельник: значение}, число_коллизий, последняя дата наблюдения)."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    pts, collisions, last_obs = {}, 0, None
    for ws in wb.worksheets:
        if not re.fullmatch(r"\d{4}", ws.title.strip()):
            continue
        year = int(ws.title)
        rows = list(ws.iter_rows(values_only=True))
        hdr = next((r for r in rows if norm(r[0]).lower() == "наименование"), None)
        row = next((r for r in rows if norm(r[0]).lower().startswith("дизельное топливо")), None)
        if hdr is None or row is None:
            raise CollectError("Лист «{}»: не найдены заголовок дат или строка «Дизельное топливо».".format(ws.title))
        for h, v in zip(hdr[1:], row[1:]):
            m = re.fullmatch(r"на (\d{1,2}) ([а-я]+)", norm(h).lower())
            val = cell_value(v)
            if not m or m.group(2) not in MONTHS_GEN or val is None:
                continue
            d = date(year, MONTHS_GEN.index(m.group(2)) + 1, int(m.group(1)))
            monday = (d - timedelta(days=d.weekday())).isoformat()
            if monday in pts:
                collisions += 1  # две даты в одной неделе (напр. 10 и 14 января 2022): остаётся более поздняя
            pts[monday] = val
            last_obs = max(last_obs or d, d)
    if not pts:
        raise CollectError("В файле {} не найдена цена дизтоплива.".format(path.name))
    return dict(sorted(pts.items())), collisions, last_obs


# ── запись рядов ──────────────────────────────────────────────────────────────
CROP_NAME = {"wheat": "пшеница", "barley": "ячмень", "sunflower": "семена подсолнечника", "corn": "кукуруза"}


def build(cena_path, nedel_path):
    prices, updated = parse_prices(cena_path)
    cena_url = url_for(cena_path.name, PAGE_URL)
    written = []
    for (scope, crop), pts in sorted(prices.items()):
        points = sorted(pts.items())
        first_real = next((i for i, (_, v) in enumerate(points) if v is not None), None)
        if first_real is None:
            continue
        points = points[first_real:]
        gaps = [d[:7] for d, v in points if v is None]
        where = "РФ" if scope == "russia" else "субъект РФ ({})".format(scope)
        note = ("Реальные данные Росстата, файл {}: {}, культура «{}» (ОКПД2 {}), на конец периода, руб./т. "
                "Только месяцы, опубликованные Росстатом ({} … {}); демо-ряды с ними не склеены. Пропуски («...», «-») — null: {}.").format(
            cena_path.name, where, CROP_NAME[crop], CROP_CODES[crop], points[0][0][:7], points[-1][0][:7],
            ", ".join(gaps) if gaps else "нет")
        d = make_series("producer_price", scope, crop, "руб./т", PRICE_SOURCE, cena_url,
                        updated or date.today().isoformat(), "month", False, points, note=note)
        errs = validate_series(d)
        if errs:
            raise CollectError("Ряд {} {} не прошёл проверку: {}".format(scope, crop, errs))
        written.append((write_series(d), d))

    dsr, collisions, last_obs = parse_diesel(nedel_path)
    points = list(dsr.items())
    note = ("Реальные данные Росстата, файл nedel_sred_cen.xlsx: средняя потребительская цена дизельного топлива по РФ, руб./л. "
            "Среднее по России, не по региону. Дата точки — понедельник недели наблюдения; Росстат указывает дату вида «на 26 октября» "
            "(последняя: {}), год берётся из имени листа. ".format(last_obs.isoformat()))
    if collisions:
        note += "Если в одну неделю попали две даты Росстата (напр. 10 и 14 января 2022), оставлена более поздняя ({} случай).".format(collisions)
    d = make_series("diesel_price", "russia", None, "руб./л", DIESEL_SOURCE, url_for("nedel_sred_cen.xlsx", PAGE_URL),
                    last_obs.isoformat(), "week", False, points, note=note)
    errs = validate_series(d)
    if errs:
        raise CollectError("Ряд дизеля не прошёл проверку: {}".format(errs))
    written.append((write_series(d), d))

    for rid in set(region_names().values()):  # региональные демо-ряды дизеля заменены средним по РФ
        old = DATA_DIR / "diesel_price_{}_week.json".format(rid)
        if old.exists() and json.loads(old.read_text(encoding="utf-8")).get("demo") is True:
            old.unlink()
            print("удалён демо-ряд", old.name)
    return written


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--offline", action="store_true", help="не ходить в сеть, разобрать файлы из raw/")
    a = ap.parse_args()
    try:
        cena, nedel = latest_local() if a.offline else download()
        written = build(cena, nedel)
    except CollectError as e:
        raise SystemExit("Ошибка: {}".format(e))
    m = build_manifest()
    for name, d in written:
        s = d["series"]
        nulls = sum(1 for p in s if p["value"] is None)
        print("{:48} точек {:4}  {} … {}  пропусков {}".format(name, len(s), s[0]["date"], s[-1]["date"], nulls))
    print("index.json: {} рядов".format(len(m["files"])))
    r = subprocess.run([sys.executable, str(Path(__file__).with_name("validate_data.py"))])
    sys.exit(r.returncode)


if __name__ == "__main__":
    main()
